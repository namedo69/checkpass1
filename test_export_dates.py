import io
import json
import threading
import unittest
from datetime import date, datetime
from types import SimpleNamespace
from unittest.mock import Mock

from openpyxl import load_workbook

from export_dates import registration_date


class ExportDateTests(unittest.TestCase):
    def rows(self):
        return [
            {"stt": str(i), "account": "sample", "status": "OK", "level": "12", "registerDate": value}
            for i, value in enumerate(["02/01/2026", "15/12/2025 23:59:59", "", "invalid"], 1)
        ]

    def prepare(self, handler):
        handler.wfile = io.BytesIO()
        handler.send_response = Mock()
        handler.send_header = Mock()
        handler.end_headers = Mock()
        handler._security_headers = Mock()
        handler.security_headers = Mock()
        handler._json = Mock()
        handler.send_json = Mock()
        return handler

    def check_workbook(self, handler):
        handler._json.assert_not_called()
        handler.send_json.assert_not_called()
        workbook = load_workbook(io.BytesIO(handler.wfile.getvalue()))
        sheet = workbook.active
        column = next(cell.column for cell in sheet[1] if cell.value == "Ngày tạo")
        cells = [sheet.cell(row, column) for row in range(2, 6)]
        self.assertEqual(cells[0].value, datetime(2026, 1, 2))
        self.assertEqual(cells[1].value, datetime(2025, 12, 15))
        self.assertTrue(all(cell.is_date for cell in cells[:2]))
        self.assertTrue(all(cell.number_format == "dd/mm/yyyy" for cell in cells))
        self.assertIsNone(cells[2].value)
        self.assertIsNone(cells[3].value)
        self.assertEqual(sorted(cell.value for cell in cells[:2]), [datetime(2025, 12, 15), datetime(2026, 1, 2)])

    def test_master_export_writes_real_dates(self):
        from master_server import MasterHandler
        handler = self.prepare(object.__new__(MasterHandler))
        handler.path = "/api/jobs/sample/export.xlsx?min_level=12"
        handler._check_job_access = Mock(return_value=(True, {"id": "sample"}))
        store = Mock()
        store.fetch.side_effect = [[(1, json.dumps(row)) for row in self.rows()], []]
        handler.server = SimpleNamespace(store=store)
        handler._handle_job_export_xlsx("sample", {})
        self.check_workbook(handler)

    def test_local_export_writes_real_dates(self):
        from garena_api_test_chrome1 import Handler
        handler = self.prepare(object.__new__(Handler))
        handler.path = "/api/batch/export-xlsx"
        handler.authorized = Mock(return_value=True)
        body = json.dumps({"required_level": 12}).encode()
        handler.rfile = io.BytesIO(body)
        handler.headers = {"Content-Length": str(len(body)), "X-API-Test-Token": "test"}
        handler.server = SimpleNamespace(csrf_token="test", batch_lock=threading.Lock(), batch_rows=self.rows())
        handler.do_POST()
        self.check_workbook(handler)

    def test_parser_supports_dates_and_legacy_values(self):
        for value in ("15/08/2026", "15/08/2026 01:17:46", "2026-08-15", date(2026, 8, 15), datetime(2026, 8, 15, 1, 17)):
            self.assertEqual(registration_date(value), date(2026, 8, 15))
        for value in (None, "", "31/02/2026", "invalid"):
            self.assertIsNone(registration_date(value))


if __name__ == "__main__":
    unittest.main()
